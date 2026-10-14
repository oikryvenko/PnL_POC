from __future__ import annotations

import csv
import hashlib
import importlib.util
import shutil
from pathlib import Path

import pytest
from openpyxl import load_workbook
from pptx import Presentation


ROOT = Path(__file__).resolve().parents[1]
PROJECT_DIR = ROOT / "docs" / "project-management"
MASTER = PROJECT_DIR / "Trading_Project_Master.xlsx"
GENERATOR_PATH = PROJECT_DIR / "generate_deliverables.py"
SPEC = importlib.util.spec_from_file_location("project_management_generator_test", GENERATOR_PATH)
GENERATOR = importlib.util.module_from_spec(SPEC)
assert SPEC and SPEC.loader
SPEC.loader.exec_module(GENERATOR)


def copy_master(tmp_path: Path) -> Path:
    copied = tmp_path / "Trading_Project_Master.xlsx"
    shutil.copy2(MASTER, copied)
    return copied


def row_for(sheet, key_column: str, key: str) -> int:
    headers = [cell.value for cell in sheet[1]]
    key_index = headers.index(key_column) + 1
    for row_number in range(2, sheet.max_row + 1):
        if sheet.cell(row_number, key_index).value == key:
            return row_number
    raise AssertionError(f"Could not find {key_column}={key} in {sheet.title}")


def set_cell_by_headers(sheet, row_number: int, column_name: str, value) -> None:
    headers = [cell.value for cell in sheet[1]]
    sheet.cell(row_number, headers.index(column_name) + 1).value = value


def test_load_master_workbook_returns_existing_project_records() -> None:
    model, tests, activities = GENERATOR.load_master_model(MASTER)
    assert len(model["business_requirements"]) == 8
    assert len(model["functional_requirements"]) == 13
    assert len(model["epics"]) == 7
    assert len(model["stories"]) == 24
    assert len(model["subtasks"]) == 48
    assert len(model["data_dictionary"]) == 138
    assert len(tests) == 12
    assert len(activities) == 11


def test_missing_worksheet_fails_clearly(tmp_path: Path) -> None:
    path = copy_master(tmp_path)
    workbook = load_workbook(path)
    workbook.remove(workbook["Requirements"])
    workbook.save(path)
    with pytest.raises(ValueError, match="missing required worksheets"):
        GENERATOR.load_master_model(path)


def test_missing_required_column_fails_clearly(tmp_path: Path) -> None:
    path = copy_master(tmp_path)
    workbook = load_workbook(path)
    workbook["Requirements"]["B1"] = ""
    workbook.save(path)
    with pytest.raises(ValueError, match="Requirements: row 1"):
        GENERATOR.load_master_model(path)


def test_duplicate_requirement_id_is_rejected(tmp_path: Path) -> None:
    path = copy_master(tmp_path)
    workbook = load_workbook(path)
    sheet = workbook["Requirements"]
    first_br = row_for(sheet, "Requirement ID", "BR-001")
    second_br = row_for(sheet, "Requirement ID", "BR-002")
    set_cell_by_headers(sheet, second_br, "Requirement ID", sheet.cell(first_br, 1).value)
    workbook.save(path)
    with pytest.raises(ValueError, match="duplicate ID 'BR-001'"):
        GENERATOR.load_master_model(path)


def test_invalid_dependency_reference_is_rejected(tmp_path: Path) -> None:
    path = copy_master(tmp_path)
    workbook = load_workbook(path)
    sheet = workbook["Jira_Backlog"]
    story_row = row_for(sheet, "Issue ID", "TMP-002")
    set_cell_by_headers(sheet, story_row, "Dependencies", "TMP-DOES-NOT-EXIST")
    workbook.save(path)
    with pytest.raises(ValueError, match="invalid references.*dependencies"):
        GENERATOR.load_master_model(path)


def test_circular_dependencies_are_rejected(tmp_path: Path) -> None:
    path = copy_master(tmp_path)
    workbook = load_workbook(path)
    sheet = workbook["Jira_Backlog"]
    first = row_for(sheet, "Issue ID", "TMP-001")
    second = row_for(sheet, "Issue ID", "TMP-002")
    set_cell_by_headers(sheet, first, "Dependencies", "TMP-002")
    set_cell_by_headers(sheet, second, "Dependencies", "TMP-001")
    workbook.save(path)
    with pytest.raises(ValueError, match="circular dependency"):
        GENERATOR.load_master_model(path)


def test_subtask_parent_must_reference_a_story(tmp_path: Path) -> None:
    path = copy_master(tmp_path)
    workbook = load_workbook(path)
    sheet = workbook["Jira_Backlog"]
    row_number = row_for(sheet, "Issue ID", "TMP-001-ST01")
    set_cell_by_headers(sheet, row_number, "Parent ID", "TMP-999")
    workbook.save(path)
    with pytest.raises(ValueError, match="invalid Parent ID"):
        GENERATOR.load_master_model(path)


def test_invalid_story_points_are_rejected(tmp_path: Path) -> None:
    path = copy_master(tmp_path)
    workbook = load_workbook(path)
    sheet = workbook["Jira_Backlog"]
    row_number = row_for(sheet, "Issue ID", "TMP-001")
    set_cell_by_headers(sheet, row_number, "Story Points", 0)
    workbook.save(path)
    with pytest.raises(ValueError, match="Story Points must be positive"):
        GENERATOR.load_master_model(path)


def test_generation_publishes_reports_and_jira_csv(tmp_path: Path) -> None:
    master_copy = copy_master(tmp_path)
    output_dir = tmp_path / "published"
    GENERATOR.generate_reports(master_copy, output_dir)
    expected = {
        "Trading_Platform_Project_Plan.xlsx", "Trading_Platform_Architecture.pptx",
        "Trading_Platform_Delivery_Overview.pptx", "jira-backlog.csv",
        "Trading_Platform_Project_Deliverables.html", "README.md",
    }
    assert expected.issubset({path.name for path in output_dir.iterdir()})
    with (output_dir / "jira-backlog.csv").open(encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))
    assert len(rows) == 79
    assert {row["Issue Type"] for row in rows} == {"Epic", "Story", "Sub-task"}
    assert len(Presentation(output_dir / "Trading_Platform_Delivery_Overview.pptx").slides) >= 14
    assert len(load_workbook(output_dir / "Trading_Platform_Project_Plan.xlsx", read_only=True).sheetnames) == 25


def test_master_status_change_flows_to_csv_and_powerpoint_without_mutating_master(tmp_path: Path) -> None:
    master_copy = copy_master(tmp_path)
    original_master_hash = hashlib.sha256(master_copy.read_bytes()).hexdigest()
    workbook = load_workbook(master_copy)
    row_number = row_for(workbook["Jira_Backlog"], "Issue ID", "TMP-018")
    set_cell_by_headers(workbook["Jira_Backlog"], row_number, "Implementation Status", "BLOCKED")
    workbook.save(master_copy)
    edited_master_hash = hashlib.sha256(master_copy.read_bytes()).hexdigest()

    output_dir = tmp_path / "status-change-output"
    GENERATOR.generate_reports(master_copy, output_dir)

    with (output_dir / "jira-backlog.csv").open(encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))
    story = next(row for row in rows if row["Issue ID"] == "TMP-018")
    assert story["Implementation Status"] == "BLOCKED"
    delivery = Presentation(output_dir / "Trading_Platform_Delivery_Overview.pptx")
    status_slide = next(
        slide for slide in delivery.slides
        if any(shape.has_text_frame and "Project status |" in shape.text for shape in slide.shapes)
    )
    slide_text = "\n".join(shape.text for shape in status_slide.shapes if shape.has_text_frame)
    assert "TMP-018: BLOCKED" in slide_text
    assert hashlib.sha256(master_copy.read_bytes()).hexdigest() == edited_master_hash
    assert edited_master_hash != original_master_hash


def test_delivery_plan_schedule_change_flows_to_workbook_without_mutating_master(tmp_path: Path) -> None:
    master_copy = copy_master(tmp_path)
    workbook = load_workbook(master_copy)
    schedule_row = row_for(workbook["Delivery_Plan"], "Activity ID", "WS-01")
    set_cell_by_headers(workbook["Delivery_Plan"], schedule_row, "Start Week", 2)
    set_cell_by_headers(workbook["Delivery_Plan"], schedule_row, "Finish Week", 2)
    set_cell_by_headers(workbook["Delivery_Plan"], schedule_row, "Duration Weeks", 1)
    workbook.save(master_copy)
    edited_master_hash = hashlib.sha256(master_copy.read_bytes()).hexdigest()

    output_dir = tmp_path / "schedule-change-output"
    GENERATOR.generate_reports(master_copy, output_dir)

    report = load_workbook(output_dir / "Trading_Platform_Project_Plan.xlsx", read_only=True, data_only=True)
    schedule = report["Integrated Delivery Plan"]
    row_number = row_for(schedule, "Workstream ID", "WS-01")
    headers = [cell.value for cell in schedule[1]]
    assert schedule.cell(row_number, headers.index("Duration (weeks)") + 1).value == 1
    assert schedule.cell(row_number, headers.index("Start week") + 1).value == 2
    assert schedule.cell(row_number, headers.index("End week") + 1).value == 2
    assert hashlib.sha256(master_copy.read_bytes()).hexdigest() == edited_master_hash


def test_failed_generation_preserves_previous_output_directory(tmp_path: Path) -> None:
    master_copy = copy_master(tmp_path)
    workbook = load_workbook(master_copy)
    workbook.remove(workbook["Requirements"])
    workbook.save(master_copy)
    output_dir = tmp_path / "published"
    output_dir.mkdir()
    marker = output_dir / "previous-success.txt"
    marker.write_text("previous reports remain", encoding="utf-8")

    with pytest.raises(ValueError, match="missing required worksheets"):
        GENERATOR.generate_reports(master_copy, output_dir)
    assert marker.read_text(encoding="utf-8") == "previous reports remain"